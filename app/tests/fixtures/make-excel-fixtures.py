# SPDX-FileCopyrightText: 2026 SAI Bilişim
# SPDX-License-Identifier: AGPL-3.0-only
# Regenerates the Excel import fixtures used by tests/e2e/e2e-import-formats.cjs.
#   python tests/fixtures/make-excel-fixtures.py        (needs openpyxl)
#
# import-excel.xlsx       — the awkward shapes real users' files have: Turkish
#                           headers, a date column, a duplicated header, an
#                           empty cell, a number stored as text ("24"), and
#                           leading-zero codes stored as text ("001").
# import-excel-mixed.xlsx — columns whose first rows disagree with later rows
#                           (a typed read_xlsx aborts on these).
import datetime
import os
from openpyxl import Workbook

here = os.path.dirname(os.path.abspath(__file__))

wb = Workbook()
ws = wb.active
ws.title = "Veri"
ws.append(["Öğrenci No", "Cinsiyet", "Yaş", "Puan", "Doğum Tarihi", "Toplam", "Not", "Puan", "Kod"])
rows = [
    [1, "Kadın", 21, 85.5, datetime.date(2003, 4, 1), None, "iyi", 3, "001"],
    [2, "Erkek", 22, 90, datetime.date(2002, 12, 31), None, None, 4, "002"],
    [3, "Kadın", None, 77.25, datetime.date(2001, 1, 15), None, "orta", 5, "010"],
    [4, "Erkek", 24, "", None, None, "", 6, "100"],
]
for i, r in enumerate(rows, start=2):
    ws.append(r)
    ws.cell(row=i, column=6, value=f"=C{i}+D{i}")  # formula (no cached value from openpyxl)
ws["C5"].value = "24"  # a number stored as text
wb.create_sheet("Notlar").append(["ignored"])
wb.save(os.path.join(here, "import-excel.xlsx"))

wb = Workbook()
ws = wb.active
ws.append(["a_num_then_text", "b_empty_first", "c_date_then_text", "d_text_then_num"])
ws.append([1, None, datetime.date(2020, 1, 1), "x"])
ws.append([2, 5, datetime.date(2020, 1, 2), 7])
ws.append(["abc", 6, "bilinmiyor", 8])
for i in range(40):
    ws.append([i, i, datetime.date(2020, 2, 1), i])
wb.save(os.path.join(here, "import-excel-mixed.xlsx"))
